"""
健康报告自测：报告生成 + Word 导出（FR-08 / AC-13）

推荐运行方式（在项目根目录下执行）：
    python -m app.test.test_report
    python app/test/test_report.py

演示内容：
    1. mcp_server.tools.report_tool.generate_report(user_id, type) —— 纯函数，可脱离 MCP 直接调用；
       分别生成 daily（日报）与 weekly（周报）；
    2. 校验周报四章节：指标概览 / 异常与告警汇总 / 用药依从性 / 个性化建议（AC-13）；
    3. 校验建议是否带溯源标注（source / version / page，AC-14）；
    4. 演示导出 Word（python-docx）：把 Markdown 报告写成 .docx 落到 app/html/reports/；
       缺少 python-docx 时打印安装提示（不报错）。

依赖：mcp_server 包（tools-mcp 同学实现）、python-docx（可选）。
"""
import os
import sys
from datetime import datetime

# ---------- 项目根目录加入 sys.path ----------
_ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
if _ROOT not in sys.path:
    sys.path.insert(0, _ROOT)

# 报告导出目录（与 TTS 一样放在静态目录下，便于前端直接下载）
REPORT_DIR = os.path.join(_ROOT, "app", "html", "reports")

# 周报必须包含的四个章节（AC-13）
WEEKLY_SECTIONS = ["指标概览", "异常", "用药依从性", "建议"]


def check_weekly_sections(text: str) -> None:
    """逐节检查周报四章节是否齐全（AC-13）"""
    print("\n     周报四章节检查（AC-13）：")
    for sec in WEEKLY_SECTIONS:
        print(f"       {'✔' if sec in text else '✘'} {sec}")
    if all(sec in text for sec in WEEKLY_SECTIONS):
        print("     ✔ 四个章节齐全")
    else:
        print("     ✘ 有章节缺失（LLM 不可用时会退化为静态模板，请检查 report_tool 的降级实现）")

    # 溯源标注（AC-14）：建议里应出现 来源/版本/页码
    print("\n     建议溯源标注检查（AC-14）：")
    for key in ("依据", "来源", "第", "页"):
        print(f"       {'✔' if key in text else 'ℹ️'} 出现\"{key}\"")
    if "依据" not in text and "来源" not in text:
        print("     ⚠️ 建议里没看到溯源标注，请确认 knowledge_tool 返回了 source/version/page")


def export_docx(text: str, path: str, title: str) -> bool:
    """
    把 Markdown 报告导出为 Word（python-docx）。
    简化实现：标题行 → Heading，表格行 → 表格，其它行 → 段落。
    """
    try:
        from docx import Document
        from docx.shared import Pt
    except ImportError:
        print("\n     [跳过] 未安装 python-docx，无法导出 Word。安装：pip install python-docx")
        return False

    doc = Document()
    doc.add_heading(title, level=0)

    lines = text.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()

        # 表格：连续的 | 开头行
        if line.startswith("|"):
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                block.append(lines[i].strip())
                i += 1
            rows = []
            for b in block:
                # 跳过 Markdown 表格分隔行（如 | :--- | :--- |）
                stripped = b.replace("|", "").replace(":", "").replace("-", "").strip()
                if stripped == "":
                    continue
                cells = [c.strip() for c in b.strip("|").split("|")]
                rows.append(cells)
            if rows:
                table = doc.add_table(rows=0, cols=len(rows[0]))
                table.style = "Light Grid Accent 1"
                for r in rows:
                    cells = table.add_row().cells
                    for j, val in enumerate(r[:len(rows[0])]):
                        cells[j].text = val
            continue

        if line.startswith("### "):
            doc.add_heading(line[4:], level=3)
        elif line.startswith("## "):
            doc.add_heading(line[3:], level=2)
        elif line.startswith("# "):
            doc.add_heading(line[2:], level=1)
        elif line.startswith("- ") or line.startswith("* "):
            doc.add_paragraph(line[2:], style="List Bullet")
        elif line.startswith("> "):
            p = doc.add_paragraph(line[2:])
            if p.runs:
                p.runs[0].italic = True
                p.runs[0].font.size = Pt(9)
        elif line.strip() in ("---", "***"):
            doc.add_paragraph("—" * 30)
        elif line.strip():
            doc.add_paragraph(line)
        i += 1

    os.makedirs(os.path.dirname(path), exist_ok=True)
    doc.save(path)
    return True


def main() -> None:
    print("=" * 72)
    print("健康报告自测：报告生成 + Word 导出（FR-08 / AC-13）")
    print("=" * 72)

    # ---------------- 1. 导入报告工具 ----------------
    try:
        from mcp_server.tools.report_tool import generate_report
    except Exception as e:
        print(f"[跳过] mcp_server.tools.report_tool 不可用：{e}")
        print("       请先完成 tools-mcp 模块（契约第 21 节），再运行本脚本。")
        print("       也可以直接调本地降级实现：")
        print("           from app.ai.tool.mcp_tool import call_mcp_tool")
        print("           await call_mcp_tool('generate_report', {'user_id':'u001','type':'weekly'})")
        return

    user_id = "u001"

    # ---------------- 2. 日报 ----------------
    print(f"\n【1】generate_report('{user_id}', type='daily')")
    daily = ""
    try:
        daily = generate_report(user_id, type="daily")
        print(f"     返回长度：{len(daily)} 字符")
        print("     ---------- 日报前 24 行 ----------")
        for line in daily.splitlines()[:24]:
            print("     | " + line)
    except Exception as e:
        print(f"     [降级] 日报生成失败：{type(e).__name__}: {e}")
        print("           预期行为：LLM 不可用时返回静态模板 + 免责声明，不抛异常。")

    # ---------------- 3. 周报 + 章节校验 ----------------
    print("\n" + "=" * 72)
    print(f"【2】generate_report('{user_id}', type='weekly') 并校验四章节（AC-13）")
    print("=" * 72)
    weekly = ""
    try:
        weekly = generate_report(user_id, type="weekly")
        print(f"     返回长度：{len(weekly)} 字符")
        print("     ---------- 周报前 24 行 ----------")
        for line in weekly.splitlines()[:24]:
            print("     | " + line)
        check_weekly_sections(weekly)
    except Exception as e:
        print(f"     [降级] 周报生成失败：{type(e).__name__}: {e}")

    # ---------------- 4. Word 导出 ----------------
    print("\n" + "=" * 72)
    print("【3】导出 Word（python-docx）")
    print("=" * 72)
    text = weekly or daily
    if not text:
        print("     [跳过] 没有报告内容可导出")
    else:
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        path = os.path.join(REPORT_DIR, f"health_report_u001_{stamp}.docx")
        ok = export_docx(text, path, "智能健康管家 · 健康报告")
        if ok:
            size = os.path.getsize(path)
            print(f"     ✔ 已导出：{path}（{size} 字节）")
            print("       生产实现：GET /report/download?user_id=u001&type=weekly")
            print("                 由 report_router 用 python-docx 生成后返回 FileResponse。")
        else:
            print("     ℹ️ 导出被跳过（缺 python-docx），报告正文仍可通过接口返回。")

    print("\n" + "=" * 72)
    print("健康报告自测结束。要点回顾：")
    print("  · 报告由 Skill（daily_report / weekly_report）驱动，章节结构固定")
    print("  · 周报四章节：指标概览 / 异常与告警汇总 / 用药依从性 / 个性化建议（AC-13）")
    print("  · 每条建议带 source/version/page 标注，可溯源（AC-14）")
    print("  · 导出 Word 走 python-docx，前端用 fetch + blob 下载（带 Authorization 头）")
    print("=" * 72)


if __name__ == "__main__":
    main()
